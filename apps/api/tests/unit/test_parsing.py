from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from budgetlens.adapters.parsing import ParseLimits, detect_csv_delimiter, parse_workbook
from budgetlens.domain.errors import PayloadTooLargeError, ValidationError

ROOT = Path(__file__).resolve().parents[4]
SAMPLE = ROOT / "sample-data"


def test_csv_accepts_utf8_bom_and_plain_utf8() -> None:
    body = "period,account_code\n2026-01,6100\n"
    plain = parse_workbook("rows.csv", body.encode("utf-8"), media_type="text/csv")
    bom = parse_workbook("rows.csv", body.encode("utf-8-sig"), media_type="text/csv")
    assert plain.headers == ["period", "account_code"]
    assert bom.headers == ["period", "account_code"]
    assert plain.delimiter == ","
    assert plain.delimiter_ambiguous is False


def test_csv_detects_semicolon_and_tab() -> None:
    semicolon = parse_workbook(
        "rows.csv", b"period;account_code\n2026-01;6100\n", media_type="text/csv"
    )
    tab = parse_workbook(
        "rows.csv", b"period\taccount_code\n2026-01\t6100\n", media_type="text/csv"
    )
    assert semicolon.delimiter == ";"
    assert tab.delimiter == "\t"


def test_ambiguous_delimiter_requires_confirmation() -> None:
    text = "period,account;department\n2026-01,6100;OPS\n"
    detected, ambiguous = detect_csv_delimiter(text)
    assert ambiguous is True
    confirmed, still_ambiguous = detect_csv_delimiter(text, confirmed=",")
    assert confirmed == ","
    assert still_ambiguous is False
    with pytest.raises(ValidationError) as exc:
        detect_csv_delimiter(text, confirmed="|")
    assert exc.value.code == "UNSUPPORTED_FILE"


def test_xlsx_converts_serial_dates_using_workbook_epoch() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["period", "amount"])
    sheet["A2"] = datetime(2026, 1, 15)
    sheet["A2"].number_format = "YYYY-MM-DD"
    sheet["B2"] = 10
    buffer = io.BytesIO()
    workbook.save(buffer)
    table = parse_workbook(
        "dates.xlsx",
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert table.rows[0].values["period"] == date(2026, 1, 15).isoformat()
    assert table.sheet_name == sheet.title
    assert sheet.title in table.available_sheets


def test_xlsx_selects_named_visible_sheet() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "HiddenSource"
    first.append(["period"])
    second = workbook.create_sheet("Actuals")
    second.append(["period", "amount"])
    second.append(["2026-02", "5"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    default = parse_workbook("book.xlsx", content)
    selected = parse_workbook("book.xlsx", content, sheet_name="Actuals")
    assert default.sheet_name == "HiddenSource"
    assert selected.sheet_name == "Actuals"
    assert selected.rows[0].values["amount"] == "5"
    with pytest.raises(ValidationError) as exc:
        parse_workbook("book.xlsx", content, sheet_name="Missing")
    assert exc.value.code == "UNSUPPORTED_FILE"


def test_xlsx_formula_is_flagged_without_evaluating() -> None:
    content = (SAMPLE / "formula.xlsx").read_bytes()
    table = parse_workbook("formula.xlsx", content)
    assert table.rows[0].formula_fields == ("amount",)
    assert table.rows[0].values["amount"].startswith("=")


def test_macro_workbook_is_rejected() -> None:
    with pytest.raises(ValidationError) as exc:
        parse_workbook("book.xlsm", b"PK\x03\x04unused")
    assert exc.value.code == "UNSUPPORTED_FILE"


def test_parse_limits_are_configurable() -> None:
    content = b"period,account_code\n2026-01,6100\n2026-02,6200\n"
    with pytest.raises(PayloadTooLargeError):
        parse_workbook("rows.csv", content, media_type="text/csv", limits=ParseLimits(max_rows=1))

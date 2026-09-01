from __future__ import annotations

import csv
import io
import zipfile
from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from budgetlens.domain.errors import PayloadTooLargeError, ValidationError, field_issue
from budgetlens.domain.importing import ParsedCellRow, WorkbookTable

CSV_MEDIA_TYPES = frozenset(
    {
        "text/csv",
        "text/plain",
        "application/csv",
        "application/vnd.ms-excel",
    }
)
XLSX_MEDIA_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/zip",
        "application/octet-stream",
    }
)
MAX_SHEETS = 8
MAX_ROWS = 100_000
MAX_COLUMNS = 40
MAX_CELLS = 500_000
MAX_CELL_LENGTH = 500
MAX_UNCOMPRESSED_BYTES = 80_000_000
MAX_COMPRESSION_RATIO = 20
DANGEROUS_ZIP_NAMES = (
    "xl/vbaProject.bin",
    "xl/externalLinks/",
    "EncryptedPackage",
    "EncryptionInfo",
)


class OpenpyxlWorkbookParser:
    def parse(
        self, filename: str, content: bytes, *, media_type: str | None = None
    ) -> WorkbookTable:
        return parse_workbook(filename, content, media_type=media_type)


def detect_kind(filename: str, content: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".xlsm") or lower.endswith(".xls") or lower.endswith(".xlsb"):
        raise ValidationError(
            "UNSUPPORTED_FILE",
            "El formato de archivo no está permitido.",
            field_errors=[
                field_issue("file", "UNSUPPORTED_FILE", "Usa un CSV UTF-8 o un XLSX sin macros.")
            ],
        )
    if lower.endswith(".xlsx"):
        if not content.startswith(b"PK"):
            raise ValidationError(
                "UNSUPPORTED_FILE",
                "La firma del archivo no corresponde a un XLSX.",
                field_errors=[
                    field_issue("file", "UNSUPPORTED_FILE", "El archivo XLSX no es válido.")
                ],
            )
        return "xlsx"
    if lower.endswith(".csv"):
        if b"\x00" in content[:4096]:
            raise ValidationError(
                "UNSUPPORTED_FILE",
                "El archivo CSV contiene datos binarios.",
                field_errors=[
                    field_issue("file", "UNSUPPORTED_FILE", "El CSV debe ser texto UTF-8.")
                ],
            )
        return "csv"
    raise ValidationError(
        "UNSUPPORTED_FILE",
        "El formato de archivo no está permitido.",
        field_errors=[
            field_issue("file", "UNSUPPORTED_FILE", "Usa un CSV UTF-8 o un XLSX sin macros.")
        ],
    )


def validate_declared_media(*, kind: str, media_type: str | None) -> None:
    if not media_type:
        return
    cleaned = media_type.split(";")[0].strip().lower()
    allowed = XLSX_MEDIA_TYPES if kind == "xlsx" else CSV_MEDIA_TYPES
    if cleaned not in allowed:
        raise ValidationError(
            "UNSUPPORTED_FILE",
            "El tipo de archivo no coincide con la extensión.",
            field_errors=[
                field_issue("file", "UNSUPPORTED_FILE", "El tipo MIME no está permitido.")
            ],
        )


def parse_workbook(
    filename: str, content: bytes, *, media_type: str | None = None
) -> WorkbookTable:
    kind = detect_kind(filename, content)
    validate_declared_media(kind=kind, media_type=media_type)
    if kind == "csv":
        return _parse_csv(content)
    return _parse_xlsx(content)


def _parse_csv(content: bytes) -> WorkbookTable:
    text = content.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = list(reader)
    if not raw_rows:
        raise ValidationError(
            "MISSING_COLUMN",
            "El archivo no tiene encabezados.",
            field_errors=[
                field_issue("file", "MISSING_COLUMN", "La primera fila debe ser el encabezado.")
            ],
        )
    headers = [_bounded_text(cell.strip()) for cell in raw_rows[0]]
    if not any(headers):
        raise ValidationError(
            "MISSING_COLUMN",
            "El archivo no tiene encabezados.",
            field_errors=[
                field_issue("file", "MISSING_COLUMN", "La primera fila debe ser el encabezado.")
            ],
        )
    rows: list[ParsedCellRow] = []
    for index, raw in enumerate(raw_rows[1:], start=2):
        if len(rows) >= MAX_ROWS:
            raise PayloadTooLargeError("El archivo excede el número máximo de filas.")
        values = {
            headers[i]: _bounded_text(raw[i] if i < len(raw) else "") for i in range(len(headers))
        }
        rows.append(ParsedCellRow(row_number=index, values=values, formula_fields=()))
    return WorkbookTable(headers=headers, rows=rows, sheet_name="csv", delimiter=delimiter)


def _reject_dangerous_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            _reject_zip_bomb(archive)
    except zipfile.BadZipFile as exc:
        raise ValidationError(
            "UNSUPPORTED_FILE",
            "El archivo XLSX no es válido.",
            field_errors=[field_issue("file", "UNSUPPORTED_FILE", "El archivo XLSX no es válido.")],
        ) from exc
    for name in names:
        lowered = name.replace("\\", "/")
        if any(item.lower() in lowered.lower() for item in DANGEROUS_ZIP_NAMES):
            raise ValidationError(
                "UNSUPPORTED_FILE",
                "El archivo contiene macros o enlaces externos.",
                field_errors=[
                    field_issue(
                        "file",
                        "UNSUPPORTED_FILE",
                        "No se aceptan macros, cifrado ni enlaces externos.",
                    )
                ],
            )


def _parse_xlsx(content: bytes) -> WorkbookTable:
    _reject_dangerous_xlsx(content)
    workbook: Workbook = load_workbook(
        filename=io.BytesIO(content),
        read_only=True,
        data_only=False,
        keep_vba=False,
        rich_text=False,
    )
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if not visible:
            raise ValidationError(
                "UNSUPPORTED_FILE",
                "El libro no tiene hojas visibles.",
                field_errors=[
                    field_issue("file", "UNSUPPORTED_FILE", "El libro no tiene hojas visibles.")
                ],
            )
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ValidationError(
                "UNSUPPORTED_FILE",
                "El libro excede el número de hojas permitido.",
                field_errors=[
                    field_issue(
                        "file", "UNSUPPORTED_FILE", "Reduce el número de hojas del archivo."
                    )
                ],
            )
        sheet = visible[0]
        return _sheet_table(sheet)
    finally:
        workbook.close()


def _sheet_table(sheet: Worksheet) -> WorkbookTable:
    headers: list[str] = []
    rows: list[ParsedCellRow] = []
    cells_seen = 0
    header_row_number = 0
    for row_index, raw_row in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS), start=1):
        cells_seen += len(raw_row)
        if cells_seen > MAX_CELLS:
            raise PayloadTooLargeError("El archivo excede el número máximo de celdas.")
        values = [_cell_text(cell) for cell in raw_row]
        if not headers:
            if not any(item.strip() for item in values):
                continue
            headers = [item.strip() for item in values]
            header_row_number = row_index
            continue
        if row_index - header_row_number > MAX_ROWS:
            raise PayloadTooLargeError("El archivo excede el número máximo de filas.")
        mapping = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        formulas = tuple(
            headers[i] for i, cell in enumerate(raw_row) if i < len(headers) and _is_formula(cell)
        )
        rows.append(ParsedCellRow(row_number=row_index, values=mapping, formula_fields=formulas))
    if not headers:
        raise ValidationError(
            "MISSING_COLUMN",
            "El archivo no tiene encabezados.",
            field_errors=[
                field_issue("file", "MISSING_COLUMN", "La primera fila debe ser el encabezado.")
            ],
        )
    return WorkbookTable(headers=headers, rows=rows, sheet_name=sheet.title, delimiter=None)


def _reject_zip_bomb(archive: zipfile.ZipFile) -> None:
    uncompressed = 0
    compressed = 0
    for info in archive.infolist():
        uncompressed += max(0, info.file_size)
        compressed += max(0, info.compress_size)
        if info.file_size > MAX_UNCOMPRESSED_BYTES:
            raise PayloadTooLargeError("El archivo comprimido excede el tamaño permitido.")
    if uncompressed > MAX_UNCOMPRESSED_BYTES:
        raise PayloadTooLargeError("El archivo comprimido excede el tamaño permitido.")
    if compressed > 0 and uncompressed / compressed > MAX_COMPRESSION_RATIO:
        raise PayloadTooLargeError("El archivo comprimido está desproporcionado.")


def _bounded_text(value: str) -> str:
    if len(value) > MAX_CELL_LENGTH:
        raise ValidationError(
            "UNSUPPORTED_FILE",
            "Una celda excede la longitud máxima permitida.",
            field_errors=[
                field_issue("file", "UNSUPPORTED_FILE", "Reduce el tamaño de las celdas.")
            ],
        )
    return value


def _cell_text(cell: object) -> str:
    value = getattr(cell, "value", None)
    if value is None:
        return ""
    return _bounded_text(str(value))


def _is_formula(cell: object) -> bool:
    if getattr(cell, "data_type", None) == "f":
        return True
    value = getattr(cell, "value", None)
    return isinstance(value, str) and value.startswith("=")


def open_binary(buffer: BinaryIO) -> bytes:
    return buffer.read()
